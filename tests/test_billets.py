from dataclasses import replace

import numpy as np
import pytest

from motion_tracker import export
from motion_tracker.billets import Station, analyze_billets, pickup_events
from motion_tracker.ingest import FrameReader
from motion_tracker.model import LOST, Track
from motion_tracker.synthetic import BilletSpec, write_billet_clip
from motion_tracker.tracking import BlobConfig, BlobDetector, track_points

FRAME = (960, 540)
S1 = Station("S1", (340, 260, 460, 340))
S2 = Station("S2", (640, 260, 760, 340))
HEATER = (30, 260, 130, 340)


@pytest.fixture(scope="module")
def scene(tmp_path_factory):
    path = tmp_path_factory.mktemp("bil") / "b.mp4"
    spec = write_billet_clip(path)
    tracks = track_points(FrameReader(str(path)), [], consumers=[BlobDetector()])
    return path, spec, tracks


def truth_billets(spec: BilletSpec, hz=2000.0) -> list[Track]:
    out = []
    for n, t0 in enumerate(spec.starts, 1):
        t = np.arange(t0, t0 + spec.total, 1 / hz)
        x = spec.x_at(t - t0)
        ok = ~np.isnan(x)
        t, x = t[ok], x[ok]
        k = len(t)
        out.append(Track(f"Billet {n}", "billet", frame=np.arange(k, dtype=np.int32), t=t, x=x, y=np.full(k, spec.y),
                         conf=np.ones(k, np.float32), flags=np.zeros(k, np.uint8)))
    return out


def test_each_billet_gets_its_own_id_with_no_clicks(scene):
    _, spec, tracks = scene
    assert [t.name for t in tracks] == ["Billet 1", "Billet 2"]
    assert all(t.kind == "billet" and t.mode == "bright_blob" for t in tracks)
    assert tracks[0].t[0] == pytest.approx(spec.starts[0], abs=1 / 30)
    assert tracks[1].t[0] == pytest.approx(spec.starts[1], abs=1 / 30)


def test_blob_position_accuracy(scene):
    _, spec, tracks = scene
    for tr, t0 in zip(tracks, spec.starts):
        truth = spec.x_at(tr.t - t0)
        good = (tr.conf > 0.9) & ~np.isnan(tr.x) & ~np.isnan(truth)  # full blob visible
        err = np.abs(tr.x[good] - truth[good])
        print(f"{tr.name}: x error mean {err.mean():.2f}px max {err.max():.2f}px")
        assert err.mean() < 0.6 and err.max() < 2.0
        assert np.abs(tr.y[good] - spec.y).max() < 0.5


def test_billet_events_match_ground_truth(scene):
    _, spec, tracks = scene
    kw = dict(frame_size=FRAME, stations=[S1, S2], exit_zone=HEATER)
    truth, got = analyze_billets(truth_billets(spec), **kw), analyze_billets(tracks, **kw)
    assert got.count == 2
    tn = {e.name: e.t for e in truth.events}
    gn = {e.name: e.t for e in got.events}
    for need in ("Billet 1 exits induction heater", "Billet 1 placed in S1", "Billet 1 leaves S1", "Billet 1 placed in S2",
                 "Billet 1 leaves S2", "Billet 1 leaves press", "Billet 1 leaves frame", "Billet 2 exits induction heater"):
        assert need in gn, need
    errs = {n: abs(gn[n] - tn[n]) for n in gn if n in tn and "frame" not in n}
    print("event time error (ms):", {n: round(v * 1000, 1) for n, v in errs.items()})
    # station crossings are sub-frame; heater exit is the frame the billet first appears
    for n, v in errs.items():
        assert v < (1.0 if "exits induction" in n else 0.5) / spec.fps, n


def test_dwell_transfer_count_and_gap(scene):
    _, spec, tracks = scene
    b = analyze_billets(tracks, FRAME, stations=[S1, S2], exit_zone=HEATER)
    assert b.count == 2 and b.gaps == pytest.approx([6.0], abs=1 / 30)  # starts 0.5 and 6.5
    assert len(b.dwell["S1"]) == 2 and len(b.dwell["S2"]) == 2
    # dwell is the time inside the station rectangle: longer than the pure stop because the billet creeps in and out
    assert all(spec.dwell1 < d < spec.dwell1 + 1.5 for d in b.dwell["S1"])
    assert len(b.transfer["S1 -> S2"]) == 2 and all(0 < d < spec.move2 for d in b.transfer["S1 -> S2"])
    head, rows = export.billet_rows(b)
    assert head[:4] == ["billet", "first_seen_s", "last_seen_s", "gap_since_previous_s"] and len(rows) == 2
    assert rows[0][3] == "" and rows[1][3] != ""
    assert "transfer_S1_to_S2_s" in head


def test_billet_hidden_for_a_moment_keeps_its_id(tmp_path):
    spec = BilletSpec(seconds=8, starts=[0.5], hide=[(1.5, 1.8)])
    p = tmp_path / "h.mp4"
    write_billet_clip(p, spec)
    (tr,) = track_points(FrameReader(str(p)), [], consumers=[BlobDetector()])
    hidden = (tr.t >= 1.55) & (tr.t < 1.8)
    assert (tr.flags[hidden] & LOST).all()  # carried through the gap, flagged, never invented
    assert np.isnan(tr.x[hidden]).all()
    assert tr.t[-1] > 5  # still the same billet afterwards


def test_roi_and_exclude_and_noise_rejection(scene):
    path, spec, _ = scene
    r = FrameReader(str(path))
    none = track_points(r, [], consumers=[BlobDetector(BlobConfig(roi=(0, 0, 960, 100)))])
    assert none == []  # nothing bright up there
    exc = track_points(r, [], consumers=[BlobDetector(BlobConfig(exclude=((0, 200, 960, 400),)))])
    assert exc == []
    big = track_points(r, [], consumers=[BlobDetector(BlobConfig(min_area=100000))])
    assert big == []


def test_picked_up_by_tool_proximity():
    n = 60
    t = np.arange(n) / 30
    mk = lambda name, x, y: Track(name, frame=np.arange(n, dtype=np.int32), t=t, x=x, y=y,  # noqa: E731
                                  conf=np.ones(n, np.float32), flags=np.zeros(n, np.uint8))
    billet = mk("Billet 1", np.full(n, 100.0), np.full(n, 100.0))
    tx = np.where((np.arange(n) >= 20) & (np.arange(n) < 40), 105.0, 400.0)
    ev = pickup_events(billet, mk("Tongs", tx, np.full(n, 100.0)), max_dist=20, min_hold=3)
    assert [e.name for e in ev] == ["Billet 1 picked up by Tongs", "Billet 1 released by Tongs"]
    assert ev[0].t == pytest.approx(20 / 30) and ev[1].t == pytest.approx(39 / 30)


def _billet_run():
    from motion_tracker.runs import Run
    return Run(name="billets", points=[], billets={
        "detect": {"threshold": 220, "min_area": 150},
        "stations": [{"name": "S1", "rect": list(S1.rect)}, {"name": "S2", "rect": list(S2.rect)}],
        "exit_zone": list(HEATER)})


def test_billet_run_and_batch(scene, tmp_path):
    from openpyxl import load_workbook
    from motion_tracker.runs import Run, apply_run, run_batch
    from motion_tracker.session import Session
    path, spec, _ = scene
    res = apply_run(_billet_run(), str(path), tmp_path / "b.mtp", plot_path=tmp_path / "b.png")
    assert res.billets.count == 2 and res.warnings == [] and (tmp_path / "b.png").stat().st_size > 5000
    with Session(tmp_path / "b.mtp") as s:
        assert [t.name for t in s.load_tracks()] == ["Billet 1", "Billet 2"]
        assert any(e.name == "Billet 2 placed in S2" for e in s.load_events())
    summary = run_batch(_billet_run(), [str(path)], tmp_path / "batch", xlsx=True)
    rows = {r[1]: r for r in load_workbook(summary).active.iter_rows(values_only=True)}
    assert rows["Billet count"][2] == 2 and rows["S1 dwell"][2] == 2 and "Gap between billets" in rows
    assert (tmp_path / "batch" / "b" / "billets.xlsx").exists()


def test_billet_run_validation_and_empty_warning(scene, tmp_path):
    from motion_tracker.runs import Run, apply_run
    bad = _billet_run()
    bad.billets["detect"]["nonsense"] = 1
    with pytest.raises(ValueError, match="bad blob settings"):
        bad.validate()
    path, _, _ = scene
    dark = _billet_run()
    dark.billets["detect"]["threshold"] = 255  # nothing is brighter than this once min_area is high
    dark.billets["detect"]["min_area"] = 100000
    assert any("no billets detected" in w for w in apply_run(dark, str(path), tmp_path / "d.mtp").warnings)
