import json

import pytest
from openpyxl import load_workbook

from motion_tracker import plots
from motion_tracker.runs import Run, apply_run, run_batch
from motion_tracker.session import Session


def _run(spec):
    return Run(
        name="ram", points=[{"name": "ram", "x": spec.x, "y": spec.y_top}],
        analyses=[{"type": "ram", "item": "ram"},
                  {"type": "line_cross", "item": "ram", "axis": "y", "value": 250, "direction": "up", "name": "Ram passes mid-stroke"}],
        segments=[{"name": "Downstroke", "start": "Ram starts downstroke", "end": "Ram reaches bottom"}])


def test_run_roundtrip_and_validation(small_clip, tmp_path):
    _, spec = small_clip
    r = _run(spec)
    r.save(tmp_path / "r.json")
    assert Run.load(tmp_path / "r.json").points == r.points
    bad = json.loads((tmp_path / "r.json").read_text())
    bad["analyses"][0]["item"] = "nope"
    (tmp_path / "bad.json").write_text(json.dumps(bad))
    with pytest.raises(ValueError, match="unknown point"):
        Run.load(tmp_path / "bad.json")
    bad["analyses"][0] = {"type": "magic", "item": "ram"}
    (tmp_path / "bad.json").write_text(json.dumps(bad))
    with pytest.raises(ValueError, match="unknown analysis"):
        Run.load(tmp_path / "bad.json")


def test_apply_run_and_batch(small_clip, tmp_path):
    clip, spec = small_clip
    res = apply_run(_run(spec), str(clip), tmp_path / "one.mtp", plot_path=tmp_path / "one.png")
    names = {e.name for e in res.events}
    assert {"Ram starts downstroke", "Ram reaches bottom", "Ram passes mid-stroke"} <= names
    assert res.segment_stats["Downstroke"].n >= 1 and res.warnings == []
    assert (tmp_path / "one.png").stat().st_size > 5000
    with Session(tmp_path / "one.mtp") as s:
        assert s.get_meta("run")["name"] == "ram" and len(s.load_events()) == len(res.events)
    summary = run_batch(_run(spec), [str(clip), str(clip)], tmp_path / "batch", xlsx=True)
    ws = load_workbook(summary).active
    assert ws.max_row > 2 and ws["A2"].value == "small"
    assert (tmp_path / "batch" / "small" / "events.xlsx").exists()


def test_batch_warns_when_seed_is_wrong(small_clip, tmp_path):
    clip, spec = small_clip
    r = _run(spec)
    r.points[0].update(x=40, y=330)  # nowhere near the dot
    res = apply_run(r, str(clip), tmp_path / "w.mtp")
    assert any("check the seed" in w or "lost" in w for w in res.warnings)


def test_plot_pdf(tracked, tmp_path):
    p = plots.plot_tracks([tracked], [], tmp_path / "p.pdf")
    assert p.read_bytes()[:4] == b"%PDF"
