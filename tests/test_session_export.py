import csv

import numpy as np
import pytest
from openpyxl import load_workbook

from motion_tracker import export
from motion_tracker.events import detect_ram
from motion_tracker.ingest import probe
from motion_tracker.model import LOST, Scale
from motion_tracker.session import Session


def test_session_roundtrip_is_reopenable(clip, tracked, tmp_path):
    tr = tracked.copy()
    tr.x[5] = tr.y[5] = np.nan
    tr.flags[5] = LOST
    ram = detect_ram(tracked)
    p = tmp_path / "s.mtp"
    with Session(p) as s:
        s.set_video(probe(str(clip[0])))
        s.save_track(tr)
        s.save_track(tr)  # saving again replaces, not duplicates
        s.save_events(ram.events)
        s.set_meta("scale", {"units_per_px": 0.5})
    with Session(p) as s:
        (got,) = s.load_tracks()
        assert len(s.load_tracks()) == 1 and s.video().fps == pytest.approx(30)
        assert np.isnan(got.x[5]) and got.flags[5] == LOST and got.params["algo"] == "csrt"
        np.testing.assert_allclose(np.nan_to_num(got.y), np.nan_to_num(tr.y))
        assert [e.name for e in s.load_events()] == [e.name for e in ram.events]
        assert s.get_meta("scale") == {"units_per_px": 0.5}


def test_track_csv_pixels_by_default_and_scaled_when_calibrated(tracked, tmp_path):
    head, rows = export.track_rows([tracked])
    assert head == ["item", "frame", "timestamp_s", "x_px", "y_px", "confidence", "status"]
    scale = Scale.from_points((0, 0), (100, 0), 50.0, "mm")
    assert scale.units_per_px == pytest.approx(0.5)
    head2, rows2 = export.track_rows([tracked], scale)
    assert "x_mm" in head2 and float(rows2[3][head2.index("y_mm")]) == pytest.approx(float(rows2[3][4]) * 0.5, abs=1e-3)
    out = export.write_table(tmp_path / "t.csv", head, rows)
    with open(out) as f:
        assert len(list(csv.reader(f))) == len(tracked) + 1


def test_lost_and_interpolated_are_labelled_in_export(tracked):
    tr = tracked.copy()
    tr.x[3] = tr.y[3] = np.nan
    tr.flags[3] = LOST
    _, rows = export.track_rows([tr])
    assert rows[3][-1] == "lost" and rows[3][3] == "" and rows[2][-1] == "measured"


def test_event_log_and_summary_xlsx(tracked, tmp_path):
    ram = detect_ram(tracked)
    head, rows = export.event_rows(ram.events)
    assert rows[0][-1] == "" and rows[1][-1] != ""
    p = export.write_table(tmp_path / "events.xlsx", head, rows)
    assert load_workbook(p).active.max_row == len(rows) + 1
    sh, sr = export.summary_rows(ram)
    p = export.write_table(tmp_path / "sum.xlsx", sh, sr)
    ws = load_workbook(p).active
    assert ws["A2"].value.startswith("Cycle time") and ws["B2"].value == 2
